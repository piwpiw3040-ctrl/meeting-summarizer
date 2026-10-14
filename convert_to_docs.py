import os
import re
import sys
import subprocess
from pathlib import Path

# Ensure console supports UTF-8 on Windows
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Font family
FONT_NAME = "TH Sarabun New"
FONT_FALLBACK = "Angsana New"

# Palette
COLOR_NAVY = RGBColor(27, 54, 93)      # #1B365D
COLOR_BLUE = RGBColor(43, 108, 176)    # #2B6CB0
COLOR_DARK = RGBColor(45, 55, 72)      # #2D3748
COLOR_MUTED = RGBColor(113, 128, 150)  # #718096
HEX_NAVY = "1B365D"
HEX_LIGHT_BG = "F7FAFC"
HEX_BORDER = "CBD5E0"
HEX_ALT_ROW = "F1F5F9"
HEX_CALLOUT_BG = "EDF2F7"


def set_run_font(run, font_name=FONT_NAME, size_pt=13, bold=False, italic=False, color=COLOR_DARK):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

    # Set complex script (Thai) font
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(
        f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>'
    )
    rPr.append(rFonts)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner cell padding in dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_cell_background(cell, hex_color):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_table_borders(table, color=HEX_BORDER, sz="4"):
    """Apply clean light borders to table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def add_formatted_text(paragraph, text, base_font=FONT_NAME, base_size=13, base_color=COLOR_DARK):
    """Parse inline **bold** and *italic* markdown tags."""
    # Pattern to match **bold** or *italic*
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, font_name=base_font, size_pt=base_size, bold=True, color=base_color)
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, font_name=base_font, size_pt=base_size, italic=True, color=base_color)
        else:
            run = paragraph.add_run(token)
            set_run_font(run, font_name=base_font, size_pt=base_size, bold=False, color=base_color)


def markdown_to_docx(md_path: Path, docx_path: Path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()

    # Set page layout to A4, 2.0 cm margins
    for section in doc.sections:
        section.page_width = Inches(8.27)   # A4 width
        section.page_height = Inches(11.69) # A4 height
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

        # Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("บันทึกการประชุมภายใน (Executive Meeting Minutes)")
        set_run_font(hrun, size_pt=10, color=COLOR_MUTED)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("เอกสารสรุปการประชุม — ระบบอัตโนมัติ Meeting Summarizer")
        set_run_font(frun, size_pt=9, color=COLOR_MUTED)

    in_table = False
    table_lines = []
    in_executive_summary = False

    i = 0
    while i < len(lines):
        line = lines[i].strip()

        # Table detection
        if line.startswith('|') and line.endswith('|'):
            table_lines.append(line)
            i += 1
            # Check if table continues
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Process table
            render_table(doc, table_lines)
            table_lines = []
            continue

        # Skip horizontal dividers
        if line == '---' or line == '***':
            i += 1
            continue

        # Skip empty lines
        if not line:
            i += 1
            continue

        # Title: # ...
        if line.startswith('# '):
            title_text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(12)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_formatted_text(p, title_text, base_size=20, base_color=COLOR_NAVY)
            i += 1
            continue

        # Section Heading: ## ...
        if line.startswith('## '):
            heading_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            add_formatted_text(p, heading_text, base_size=15, base_color=COLOR_NAVY)
            i += 1
            continue

        # Sub-heading: ### ...
        if line.startswith('### '):
            heading_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            add_formatted_text(p, heading_text, base_size=13.5, base_color=COLOR_BLUE)
            i += 1
            continue

        # Nested Bullet list item: indentation + - ...
        if lines[i].startswith('    - ') or lines[i].startswith('\t- ') or lines[i].startswith('  - '):
            p = doc.add_paragraph(style='List Bullet 2' if 'List Bullet 2' in doc.styles else 'List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            bullet_text = lines[i].strip()[2:].strip()
            add_formatted_text(p, bullet_text, base_size=12.5)
            i += 1
            continue

        # Standard Bullet list item: - ...
        if line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            bullet_text = line[2:].strip()
            add_formatted_text(p, bullet_text, base_size=13)
            i += 1
            continue

        # Numbered list: 1. ...
        if re.match(r'^\d+\.\s', line):
            p = doc.add_paragraph(style='List Number' if 'List Number' in doc.styles else None)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_text(p, line, base_size=13)
            i += 1
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_text(p, line, base_size=13)
        i += 1

    doc.save(str(docx_path))
    print(f"✅ สร้างเอกสาร Word สำเร็จ: {docx_path.name}")


def render_table(doc, table_lines):
    if len(table_lines) < 2:
        return

    # Parse header
    raw_header = [c.strip() for c in table_lines[0].split('|')[1:-1]]
    
    # Parse data rows (skip delimiter row line 1)
    rows_data = []
    for t_line in table_lines[2:]:
        cols = [c.strip() for c in t_line.split('|')[1:-1]]
        if cols:
            rows_data.append(cols)

    table = doc.add_table(rows=len(rows_data) + 1, cols=len(raw_header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    # Format header row
    hdr_cells = table.rows[0].cells
    for col_idx, heading in enumerate(raw_header):
        cell = hdr_cells[col_idx]
        set_cell_background(cell, HEX_NAVY)
        set_cell_margins(cell, top=140, bottom=140, left=150, right=150)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
        add_formatted_text(p, heading, base_size=12.5, base_color=RGBColor(255, 255, 255))

    # Make header repeat across pages
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    # Format data rows
    for row_idx, data_row in enumerate(rows_data):
        row = table.rows[row_idx + 1]
        
        # Cant split row across pages
        r_trPr = row._tr.get_or_add_trPr()
        r_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        # Alternating background
        row_bg = HEX_ALT_ROW if row_idx % 2 == 1 else "FFFFFF"

        for col_idx in range(len(raw_header)):
            cell = row.cells[col_idx]
            if row_bg != "FFFFFF":
                set_cell_background(cell, row_bg)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            val = data_row[col_idx] if col_idx < len(data_row) else ""
            p = cell.paragraphs[0]
            # Center-align index, owner, deadline
            if col_idx in [0, 2, 3]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            add_formatted_text(p, val, base_size=12)

    # Add space after table
    sp_p = doc.add_paragraph()
    sp_p.paragraph_format.space_before = Pt(6)
    sp_p.paragraph_format.space_after = Pt(6)


def convert_docx_to_pdf(docx_path: Path) -> Path:
    pdf_path = docx_path.with_suffix(".pdf")
    abs_docx = docx_path.resolve()
    abs_pdf = pdf_path.resolve()

    ps_script = f"""
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    try {{
        $doc = $word.Documents.Open('{abs_docx}')
        $doc.SaveAs([ref]'{abs_pdf}', [ref]17)
        $doc.Close()
    }} finally {{
        $word.Quit()
    }}
    """
    res = subprocess.run(["powershell", "-NoProfile", "-Command", ps_script], capture_output=True, text=True)
    if res.returncode != 0:
        raise RuntimeError(f"PowerShell PDF export failed: {res.stderr}")
    print(f"✅ แปลงเป็น PDF สำเร็จ: {pdf_path.name}")
    return pdf_path


def main():
    base_dir = Path(__file__).resolve().parent
    summaries_dir = base_dir / "summaries"

    md_files = list(summaries_dir.glob("*.md"))
    if not md_files:
        print("ไม่พบไฟล์ Markdown ในโฟลเดอร์ summaries/")
        return

    for md_file in sorted(md_files):
        docx_file = md_file.with_suffix(".docx")
        print(f"\n📄 กำลังประมวลผล: {md_file.name}")
        markdown_to_docx(md_file, docx_file)
        try:
            convert_docx_to_pdf(docx_file)
        except Exception as e:
            print(f"⚠️ เกิดข้อผิดพลาดในการแปลง PDF: {e}")


if __name__ == "__main__":
    main()
