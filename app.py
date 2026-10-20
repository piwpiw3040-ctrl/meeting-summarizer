import os
import io
import re
import sys
import time
import tempfile
from pathlib import Path
from datetime import datetime
import streamlit as st
from dotenv import load_dotenv

# Ensure console supports UTF-8 on Windows
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

# Load environment
base_dir = Path(__file__).resolve().parent
env_file = base_dir / ".env"
if env_file.exists():
    load_dotenv(dotenv_path=env_file)
else:
    load_dotenv()

# Page config for mobile
st.set_page_config(
    page_title="สรุปการประชุม AI (Meeting Summarizer)",
    page_icon="🎙️",
    layout="centered",
    initial_sidebar_state="collapsed",
)

# Custom CSS for Mobile App feel
st.markdown("""
<style>
    /* Responsive container */
    .block-container {
        padding-top: 1.5rem;
        padding-bottom: 3rem;
        padding-left: 1rem;
        padding-right: 1rem;
        max-width: 720px;
    }
    /* Mobile-first buttons */
    .stButton>button {
        width: 100%;
        border-radius: 12px;
        height: 3.2rem;
        font-weight: 600;
        font-size: 1.05rem;
        background: linear-gradient(135deg, #1B365D 0%, #2B6CB0 100%);
        color: white;
        border: none;
        box-shadow: 0 4px 12px rgba(27, 54, 93, 0.2);
        transition: all 0.2s ease;
    }
    .stButton>button:hover {
        transform: translateY(-2px);
        box-shadow: 0 6px 16px rgba(27, 54, 93, 0.3);
    }
    /* Card design */
    .meeting-card {
        background: #ffffff;
        border-radius: 14px;
        padding: 1.25rem;
        margin-bottom: 1rem;
        border: 1px solid #E2E8F0;
        box-shadow: 0 2px 8px rgba(0, 0, 0, 0.04);
    }
    /* Badge */
    .badge {
        display: inline-block;
        padding: 0.25rem 0.6rem;
        font-size: 0.8rem;
        font-weight: 600;
        border-radius: 20px;
        background-color: #EBF8FF;
        color: #2B6CB0;
    }
</style>
""", unsafe_allow_html=True)

SUMMARY_PROMPT = """คุณคือเลขานุการและผู้เชี่ยวชาญด้านการบันทึกการประชุมมืออาชีพ (Senior Executive Meeting Secretary)
โปรดฟังเสียงบันทึกการประชุมที่แนบมานี้อย่างละเอียด แล้วสรุปเนื้อหาเป็นภาษาไทยที่กระชับ ชัดเจน เป็นมืออาชีพ ไม่ตกหล่นใจความสำคัญ

กรุณาจัดโครงสร้างให้อยู่ในรูปแบบ Markdown ดังต่อไปนี้อย่างเคร่งครัด:

# 📋 บันทึกการประชุม: {meeting_name}

- **📅 วันที่บันทึกสรุป:** {date_str}
- **🎙️ ไฟล์เสียงต้นฉบับ:** {filename}
- **👥 ผู้เข้าร่วมประชุม / ผู้มีบทบาทหลัก:** (สรุปรายชื่อหรือบทบาทเท่าที่สามารถระบุได้จากการสนทนา หากไม่ทราบให้ระบุตามบริบท เช่น ทีมพัฒนา, ผู้จัดการฝ่าย ฯลฯ)

---

## 🎯 สรุปภาพรวมสำหรับผู้บริหาร (Executive Summary)
(สรุปใจความสำคัญและเป้าหมายหลักของการประชุมใน 3-5 บรรทัด ครอบคลุมว่าประชุมเรื่องอะไร ทำไม และผลลัพธ์หลักคืออะไร)

---

## ⚖️ มติที่ประชุมและข้อตกลงร่วมกัน (Decisions & Key Agreements)
(สรุปการตัดสินใจขั้นสุดท้าย นโยบาย หรือข้อตกลงที่สมาชิกในที่ประชุมเห็นชอบร่วมกัน เป็นข้อๆ)
- [มติ/ข้อตกลงที่ 1]
- [มติ/ข้อตกลงที่ 2]

---

## ✅ รายการสิ่งที่ต้องทำ (Action Items)
(จัดทำเป็นตาราง Markdown ให้ละเอียด ระบุสิ่งที่ต้องทำ ผู้รับผิดชอบ และกำหนดส่งให้ชัดเจน หากไม่ระบุเดดไลน์ชัดเจนให้ใส่ 'ตามที่ตกลง' หรือ 'รอติดตาม')

| ลำดับ | สิ่งที่ต้องทำ (Action Item) | ผู้รับผิดชอบ (Owner) | กำหนดส่ง (Deadline) | หมายเหตุ / เงื่อนไข |
| :---: | :--- | :--- | :--- | :--- |
| 1 | ... | ... | ... | ... |

---

## 📝 ประเด็นสำคัญแยกตามวาระการพูดคุย (Key Discussion Points)
(สรุปเนื้อหารายละเอียดของการพูดคุย โดยจัดกลุ่มแยกตามหัวข้อหรือวาระที่ถกกันอย่างเป็นระบบ)

### 1. [หัวข้อวาระที่ 1]
- **สาระสำคัญ:** ...
- **ความคิดเห็น/ข้อเสนอแนะ:** ...

### 2. [หัวข้อวาระที่ 2]
- **สาระสำคัญ:** ...
- **ความคิดเห็น/ข้อเสนอแนะ:** ...

---

## ⚠️ ปัญหา อุปสรรค และความเสี่ยง (Blockers & Risks)
(รวบรวมประเด็นที่เป็นปัญหา ความท้าทาย สิ่งที่กำลังติดขัด หรือความเสี่ยงที่ถูกหยิบยกขึ้นมาในที่ประชุม)
- ...

---

## 📅 การติดตามงานและนัดหมายครั้งถัดไป (Next Steps & Follow-up)
(การนัดหมายประชุมรอบถัดไป วันที่ เวลา หรือขั้นตอนการติดตามความคืบหน้า)
- ...
"""


def get_mime_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    mime_map = {
        ".mp3": "audio/mp3",
        ".m4a": "audio/mp4",
        ".wav": "audio/wav",
        ".aac": "audio/aac",
        ".ogg": "audio/ogg",
        ".flac": "audio/flac",
        ".webm": "audio/webm",
        ".mp4": "video/mp4",
        ".m4v": "video/mp4",
    }
    return mime_map.get(ext, "audio/mp3")


def create_docx_bytes(markdown_text: str) -> bytes:
    """Generate Word .docx in memory from markdown."""
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    font_name = "TH Sarabun New"
    doc = docx.Document()

    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("บันทึกการประชุมภายใน (Meeting Minutes)")
        hrun.font.name = font_name
        hrun.font.size = Pt(10)
        hrun.font.color.rgb = RGBColor(113, 128, 150)

    lines = markdown_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or line in ['---', '***']:
            i += 1
            continue

        if line.startswith('|') and line.endswith('|'):
            t_lines = [line]
            i += 1
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                t_lines.append(lines[i].strip())
                i += 1
            if len(t_lines) >= 3:
                raw_hdr = [c.strip() for c in t_lines[0].split('|')[1:-1]]
                rows_data = [[c.strip() for c in l.split('|')[1:-1]] for l in t_lines[2:]]
                tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(raw_hdr))
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

                # Header
                for idx, h in enumerate(raw_hdr):
                    c = tbl.rows[0].cells[idx]
                    tcPr = c._tc.get_or_add_tcPr()
                    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="1B365D"/>'))
                    p = c.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
                    r = p.add_run(h)
                    r.font.name = font_name
                    r.font.size = Pt(12.5)
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)

                # Rows
                for r_idx, r_data in enumerate(rows_data):
                    r_el = tbl.rows[r_idx + 1]
                    bg = "F1F5F9" if r_idx % 2 == 1 else "FFFFFF"
                    for col_idx in range(len(raw_hdr)):
                        c = r_el.cells[col_idx]
                        if bg != "FFFFFF":
                            c._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}"/>'))
                        val = r_data[col_idx] if col_idx < len(r_data) else ""
                        p = c.paragraphs[0]
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 2, 3] else WD_ALIGN_PARAGRAPH.LEFT
                        r = p.add_run(val)
                        r.font.name = font_name
                        r.font.size = Pt(12)
            continue

        p = doc.add_paragraph()
        if line.startswith('# '):
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(12)
            r = p.add_run(line[2:].strip())
            r.font.name = font_name
            r.font.size = Pt(20)
            r.font.bold = True
            r.font.color.rgb = RGBColor(27, 54, 93)
        elif line.startswith('## '):
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(line[3:].strip())
            r.font.name = font_name
            r.font.size = Pt(15)
            r.font.bold = True
            r.font.color.rgb = RGBColor(27, 54, 93)
        elif line.startswith('### '):
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(line[4:].strip())
            r.font.name = font_name
            r.font.size = Pt(13.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(43, 108, 176)
        else:
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            # Basic bold formatting
            tokens = re.split(r'(\*\*.*?\*\*)', line)
            for t in tokens:
                if t.startswith('**') and t.endswith('**'):
                    r = p.add_run(t[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(t)
                r.font.name = font_name
                r.font.size = Pt(13)
        i += 1

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def create_print_html(markdown_text: str, title: str = "บันทึกการประชุม") -> str:
    """Generate printable HTML view optimized for Mobile PDF printing/saving."""
    import html
    # Simple markdown conversion for print
    content = markdown_text
    # Bold
    content = re.sub(r'\*\*(.*?)\*\*', r'<strong>\1</strong>', content)
    # Headings
    content = re.sub(r'^# (.*?)$', r'<h1>\1</h1>', content, flags=re.MULTILINE)
    content = re.sub(r'^## (.*?)$', r'<h2>\1</h2>', content, flags=re.MULTILINE)
    content = re.sub(r'^### (.*?)$', r'<h3>\1</h3>', content, flags=re.MULTILINE)
    # Bullets
    content = re.sub(r'^- (.*?)$', r'<li>\1</li>', content, flags=re.MULTILINE)

    html_template = f"""<!DOCTYPE html>
<html lang="th">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Sarabun:ital,wght@0,300;0,400;0,600;0,700;1,400&display=swap" rel="stylesheet">
<style>
    body {{
        font-family: 'Sarabun', -apple-system, BlinkMacSystemFont, sans-serif;
        line-height: 1.6;
        color: #2D3748;
        padding: 20px;
        max-width: 800px;
        margin: 0 auto;
        background: #fff;
    }}
    h1 {{ color: #1B365D; border-bottom: 2px solid #2B6CB0; padding-bottom: 8px; font-size: 1.8rem; margin-top: 0; }}
    h2 {{ color: #1B365D; font-size: 1.35rem; margin-top: 1.5rem; border-left: 4px solid #2B6CB0; padding-left: 8px; }}
    h3 {{ color: #2B6CB0; font-size: 1.15rem; margin-top: 1rem; }}
    table {{ width: 100%; border-collapse: collapse; margin: 15px 0; font-size: 0.92rem; }}
    th, td {{ border: 1px solid #CBD5E0; padding: 8px 10px; text-align: left; }}
    th {{ background-color: #1B365D; color: white; font-weight: 600; }}
    tr:nth-child(even) {{ background-color: #F8FAFC; }}
    li {{ margin-bottom: 4px; }}
    @media print {{
        body {{ padding: 0; font-size: 11pt; }}
        .no-print {{ display: none !important; }}
    }}
</style>
</head>
<body>
<div class="no-print" style="margin-bottom: 20px; text-align: center;">
    <button onclick="window.print()" style="background:#1B365D; color:white; border:none; padding:12px 24px; border-radius:8px; font-size:16px; font-weight:bold; cursor:pointer;">
        🖨️ พิมพ์หรือบันทึกเป็น PDF ทันที
    </button>
</div>
{content}
</body>
</html>"""
    return html_template


# ================= SECURITY / PIN LOCK =================
APP_PIN = "304000"
try:
    if hasattr(st, "secrets") and "APP_PASSWORD" in st.secrets:
        APP_PIN = str(st.secrets["APP_PASSWORD"]).strip()
except Exception:
    pass

if not st.session_state.get("authenticated", False):
    st.markdown("""
    <div style="text-align: center; margin-top: 2rem; margin-bottom: 2rem;">
        <div style="font-size: 3.5rem; margin-bottom: 0.5rem;">🔒</div>
        <h2 style="color: #1B365D; margin-bottom: 0.2rem;">ระบบสรุปบันทึกการประชุม AI</h2>
        <p style="color: #718096; font-size: 0.95rem;">พื้นที่ส่วนบุคคล — กรุณาใส่รหัสผ่านเพื่อเข้าใช้งาน</p>
    </div>
    """, unsafe_allow_html=True)

    with st.form("login_form"):
        entered_pin = st.text_input("รหัสผ่าน (PIN):", type="password", placeholder="ใส่รหัสผ่าน 6 หลัก")
        login_btn = st.form_submit_button("🔓 เข้าสู่ระบบ (Unlock)")
        if login_btn:
            if entered_pin.strip() == APP_PIN:
                st.session_state["authenticated"] = True
                st.rerun()
            else:
                st.error("❌ รหัสผ่านไม่ถูกต้อง กรุณาลองใหม่อีกครั้งครับ")

    st.stop()


# ================= UI LAYOUT =================
st.markdown("""
<div style="text-align: center; margin-bottom: 1.5rem;">
    <h1 style="font-size: 1.8rem; color: #1B365D; margin-bottom: 0.2rem;">🎙️ สรุปบันทึกการประชุม AI</h1>
    <p style="color: #718096; font-size: 0.95rem;">อัปโหลดเสียงการประชุมแล้วรับบทสรุปและ Action Items ได้ทันทีบนมือถือ</p>
</div>
""", unsafe_allow_html=True)

# Read API Key from st.secrets or os.getenv
api_key_env = ""
try:
    if hasattr(st, "secrets") and "GEMINI_API_KEY" in st.secrets:
        api_key_env = str(st.secrets["GEMINI_API_KEY"]).strip()
except Exception:
    pass

if not api_key_env:
    api_key_env = os.getenv("GEMINI_API_KEY", "").strip()

# Read Model
model_default = "gemini-3.8-flash"
try:
    if hasattr(st, "secrets") and "GEMINI_MODEL" in st.secrets:
        model_default = str(st.secrets["GEMINI_MODEL"]).strip()
except Exception:
    pass
if not model_default:
    model_default = os.getenv("GEMINI_MODEL", "gemini-3.8-flash").strip()

# Sidebar
with st.sidebar:
    st.header("⚙️ การตั้งค่าระบบ")
    api_key_input = st.text_input(
        "Google Gemini API Key:",
        value=api_key_env,
        type="password",
        help="ขอ API Key ฟรีได้ที่ https://aistudio.google.com/app/apikey"
    )
    model_name = st.selectbox(
        "AI Model:",
        options=["gemini-3.8-flash", "gemini-flash-latest"],
        index=0 if model_default == "gemini-3.8-flash" else 1
    )
    if api_key_input or api_key_env:
        st.success("✅ เชื่อมต่อ API Key แล้ว")
    st.markdown("---")
    st.markdown("### 📱 วิธีติดตั้งบนมือถือ Android")
    st.markdown("""
    1. เปิดหน้านี้ในเบราว์เซอร์ **Google Chrome** บนมือถือ
    2. กดปุ่มเมนู **จุด 3 จุด (⋮)** ที่มุมขวาบน
    3. เลือก **"เพิ่มลงในหน้าจอหลัก" (Add to Home screen)**
    4. จะได้ไอคอนแอปบนหน้าจอมือถือ กดเปิดใช้ได้ตลอดเวลา!
    """)
    st.markdown("---")
    if st.button("🔒 ล็อกหน้าจอ / ออกจากระบบ (Logout)", key="logout_btn"):
        st.session_state["authenticated"] = False
        st.rerun()

# Check API key
active_key = api_key_input.strip() or api_key_env.strip()
if not active_key or active_key == "YOUR_GEMINI_API_KEY_HERE":
    st.warning("⚠️ ยังไม่พบ Gemini API Key ในระบบ")
    with st.expander("🔑 แตะที่นี่เพื่อใส่ API Key ของคุณได้ทันที", expanded=True):
        main_key_input = st.text_input(
            "วาง Gemini API Key ของคุณที่นี่:",
            type="password",
            placeholder="AIzaSy... หรือ AQ.Ab...",
            help="ขอ API Key ฟรีได้ที่ https://aistudio.google.com/app/apikey"
        )
        if main_key_input:
            st.session_state["custom_api_key"] = main_key_input.strip()
            active_key = main_key_input.strip()
            st.rerun()

    if "custom_api_key" in st.session_state:
        active_key = st.session_state["custom_api_key"]

    if not active_key or active_key == "YOUR_GEMINI_API_KEY_HERE":
        st.info("💡 **คำแนะนำ:**\n1. คุณสามารถวาง API Key ในกล่องข้อความด้านบนนี้ได้เลย หรือ\n2. ไปที่หน้าจัดการแอปบน Streamlit Cloud -> Settings -> Secrets แล้วใส่ `GEMINI_API_KEY = \"...\"`")
        st.stop()

# Tabs
tab1, tab2 = st.tabs(["🎙️ สรุปไฟล์เสียงใหม่", "📚 ประวัติการประชุมที่ผ่านมา"])

with tab1:
    st.markdown("#### 1. เลือกไฟล์เสียงการประชุม")
    audio_file = st.file_uploader(
        "อัปโหลดไฟล์เสียง (.mp3, .m4a, .wav, .aac, .ogg, .mp4):",
        type=["mp3", "m4a", "wav", "aac", "ogg", "flac", "mp4", "webm"],
        help="รองรับไฟล์เสียงจากแอปบันทึกเสียง iPhone และ Android ทุกรุ่น"
    )

    custom_title = st.text_input("หัวข้อหรือชื่อการประชุม (ไม่บังคับ):", placeholder="เช่น ประชุมวางแผนระบบ Smart Farm ไซต์งาน")

    if audio_file is not None:
        file_size_mb = audio_file.size / (1024 * 1024)
        st.caption(f"📁 ชื่อไฟล์: **{audio_file.name}** | ขนาด: **{file_size_mb:.2f} MB**")

        start_btn = st.button("🚀 เริ่มวิเคราะห์และสรุปการประชุม")

        if start_btn:
            try:
                from google import genai
                from google.genai import types

                client = genai.Client(api_key=active_key)

                # Save temporary file for upload
                suffix = Path(audio_file.name).suffix or ".mp3"
                with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp_f:
                    tmp_f.write(audio_file.getvalue())
                    tmp_path = Path(tmp_f.name)

                progress_bar = st.progress(0, text="กำลังเตรียมไฟล์เสียง...")

                # Upload to Gemini File API
                mime_type = get_mime_type(audio_file.name)
                progress_bar.progress(25, text="⬆️ กำลังอัปโหลดไฟล์เสียงไปยัง Gemini API...")
                
                with open(tmp_path, "rb") as f:
                    uploaded_remote_file = client.files.upload(
                        file=f,
                        config=types.UploadFileConfig(
                            mime_type=mime_type,
                            display_name=audio_file.name[:40]
                        )
                    )

                # Wait for active
                while uploaded_remote_file.state.name == "PROCESSING":
                    progress_bar.progress(40, text="⏳ ระบบกำลังเตรียมประมวลผลไฟล์เสียง...")
                    time.sleep(3)
                    uploaded_remote_file = client.files.get(name=uploaded_remote_file.name)

                # Build Prompt
                meeting_title = custom_title.strip() if custom_title.strip() else audio_file.name
                meeting_title = meeting_title.rsplit('.', 1)[0]
                now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
                prompt_content = SUMMARY_PROMPT.format(
                    meeting_name=meeting_title,
                    date_str=now_str,
                    filename=audio_file.name
                )

                progress_bar.progress(60, text=f"🧠 กำลังให้ AI วิเคราะห์เนื้อหาการประชุม...")

                # Call Gemini with auto-retry
                summary_text = None
                max_retries = 3
                candidate_models = [model_name]
                if "gemini-flash-latest" not in candidate_models:
                    candidate_models.append("gemini-flash-latest")

                for attempt in range(1, max_retries + 1):
                    cur_model = candidate_models[(attempt - 1) % len(candidate_models)]
                    try:
                        resp = client.models.generate_content(
                            model=cur_model,
                            contents=[uploaded_remote_file, prompt_content],
                            config=types.GenerateContentConfig(temperature=0.2)
                        )
                        summary_text = resp.text
                        if summary_text:
                            break
                    except Exception as ge:
                        err_str = str(ge)
                        if ("503" in err_str or "UNAVAILABLE" in err_str) and attempt < max_retries:
                            progress_bar.progress(75, text="⚠️ เซิร์ฟเวอร์หนาแน่น กำลังลองใหม่อัตโนมัติ...")
                            time.sleep(8)
                        else:
                            raise ge

                # Clean up cloud file
                try:
                    client.files.delete(name=uploaded_remote_file.name)
                except Exception:
                    pass

                # Clean up temp local file
                try:
                    tmp_path.unlink()
                except Exception:
                    pass

                progress_bar.progress(100, text="✅ สรุปการประชุมเรียบร้อยแล้ว!")
                time.sleep(0.5)
                progress_bar.empty()

                st.session_state["latest_summary"] = summary_text
                st.session_state["latest_title"] = meeting_title

                # Auto-save to summaries directory
                summaries_dir = base_dir / "summaries"
                summaries_dir.mkdir(exist_ok=True)
                date_prefix = datetime.now().strftime("%Y-%m-%d")
                save_stem = re.sub(r'[^\w\s-]', '', meeting_title).strip()
                out_md_path = summaries_dir / f"{date_prefix}_{save_stem}_summary.md"
                with open(out_md_path, "w", encoding="utf-8") as out_f:
                    out_f.write(summary_text)

            except Exception as e:
                st.error(f"❌ เกิดข้อผิดพลาดในการประมวลผล: {e}")

    # Display Result
    if "latest_summary" in st.session_state:
        st.markdown("---")
        st.success("🎉 ประมวลผลบทสรุปสำเร็จแล้ว!")
        
        # Download buttons
        st.markdown("### 📥 ดาวน์โหลดเอกสาร")
        dcol1, dcol2 = st.columns(2)
        
        # Word .docx
        with dcol1:
            docx_data = create_docx_bytes(st.session_state["latest_summary"])
            st.download_button(
                label="📄 โหลดไฟล์ Word (.docx)",
                data=docx_data,
                file_name=f"{st.session_state.get('latest_title', 'meeting')}_summary.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )

        # Markdown .md
        with dcol2:
            st.download_button(
                label="📝 โหลดไฟล์ Markdown (.md)",
                data=st.session_state["latest_summary"].encode("utf-8"),
                file_name=f"{st.session_state.get('latest_title', 'meeting')}_summary.md",
                mime="text/markdown",
                use_container_width=True
            )

        # Print / Save as PDF view
        print_html = create_print_html(st.session_state["latest_summary"], st.session_state.get('latest_title', 'บันทึกการประชุม'))
        st.download_button(
            label="📕 พิมพ์ / บันทึกเป็น PDF (Print view)",
            data=print_html.encode("utf-8"),
            file_name=f"{st.session_state.get('latest_title', 'meeting')}_print.html",
            mime="text/html",
            use_container_width=True,
            help="เปิดไฟล์นี้ในเบราว์เซอร์แล้วสั่ง Print/Save as PDF ได้ทันทีบนมือถือ"
        )

        st.markdown("---")
        st.markdown("### 📋 เนื้อหาบทสรุปการประชุม")
        st.markdown(st.session_state["latest_summary"])

with tab2:
    st.markdown("#### 📚 บันทึกการประชุมที่เคยประมวลผลไว้")
    summaries_dir = base_dir / "summaries"
    if summaries_dir.exists():
        md_history = sorted(list(summaries_dir.glob("*.md")), reverse=True)
        if not md_history:
            st.info("ยังไม่มีประวัติการสรุปการประชุมในระบบ")
        else:
            for md_h in md_history:
                with st.expander(f"📄 {md_h.stem}"):
                    with open(md_h, "r", encoding="utf-8") as f:
                        h_content = f.read()
                    st.markdown(h_content)
                    
                    # Direct download buttons for history
                    docx_h = create_docx_bytes(h_content)
                    st.download_button(
                        label=f"📥 ดาวน์โหลด Word ({md_h.stem}.docx)",
                        data=docx_h,
                        file_name=f"{md_h.stem}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        key=f"btn_docx_{md_h.name}"
                    )
    else:
        st.info("ยังไม่มีโฟลเดอร์ประวัติการสรุป")
